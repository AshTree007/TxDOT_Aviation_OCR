from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path


def style_sheet(sheet: object, widths: list[float], header_row: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"E{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[header_row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="234E70")
    sheet.row_dimensions[header_row].height = 42


def save_summary(summary: dict, contractors: list[dict], source_url: str, path: Path) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from pdf_extractor import clean_name

    # Follow the contractor sections in the PDF, as in the reference image.
    contractors = sorted(
        contractors,
        key=lambda contractor: min(
            (page for item in contractor["items"] for page in item.get("Pages", [])),
            default=1000000,
        ),
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bid comparison"
    project = summary["project"]
    sheet.append([f"{project['Airport name']} — Bid comparison"])
    sheet["A1"].font = Font(name="Arial", size=16, bold=True)
    sheet.append(["CSJ", project["TxDOT CSJ#"], "Scope", project["Project Scope"]])
    sheet.append(["Bid date", project["Bid Opening Date"]])
    sheet.append(["Lowest totals are apparent bids, not confirmed awards."])
    sheet.append([])

    headers = ["Item", "Qty", "Unit", "Description"]
    for contractor in contractors:
        headers.extend(["Unit Price", "Total Price"])
    sheet.append(headers)

    lookup = [
        {clean_name(item["Item"]): item for item in contractor["items"]}
        for contractor in contractors
    ]
    for item in contractors[0]["items"]:
        row = [item["Item"], float(item["Qty"]), item["Unit"], item["Description"]]
        for contractor_items in lookup:
            prices = contractor_items[clean_name(item["Item"])]
            row.extend([float(prices["Unit Price"]), float(prices["Total Price"])])
        sheet.append(row)
        sheet.row_dimensions[sheet.max_row].height = 42

    first_row, last_row = 7, sheet.max_row
    sheet.append([None])
    base_row = sheet.max_row
    sheet.append(["Calculated item totals"])
    calculated_row = sheet.max_row
    sheet.append(["Difference"])
    difference_row = sheet.max_row
    sheet.append(["Bid status"])
    lowest = min(contractor["base_bid"] for contractor in contractors)
    tied = sum(contractor["base_bid"] == lowest for contractor in contractors) > 1

    for index, contractor in enumerate(contractors):
        unit_column, total_column = 5 + index * 2, 6 + index * 2
        letter = get_column_letter(total_column)
        sheet.cell(calculated_row, total_column, f"=SUM({letter}{first_row}:{letter}{last_row})")
        sheet.cell(base_row, unit_column, "Total Base Bid:")
        sheet.cell(base_row, total_column, float(contractor["base_bid"]))
        sheet.cell(difference_row, total_column, f"={letter}{calculated_row}-{letter}{base_row}")
        if contractor["base_bid"] == lowest:
            sheet.cell(difference_row + 1, total_column, "Apparent lowest (tie)" if tied else "Apparent lowest")
            sheet.cell(base_row, total_column).fill = PatternFill("solid", fgColor="DDEFE5")
        for row in range(first_row, difference_row + 1):
            sheet.cell(row, unit_column).number_format = '"$"#,##0.00'
            sheet.cell(row, total_column).number_format = '"$"#,##0.00'
    for row in range(first_row, last_row + 1):
        quantity = sheet.cell(row, 2)
        quantity.number_format = '#,##0' if float(quantity.value).is_integer() else '#,##0.###'

    # Keep all extracted source values, descriptions and page references.
    source = workbook.create_sheet("Extracted bids")
    source.append(["Source PDF", source_url])
    source.append(["Contractor", "Item", "Qty", "Unit", "Description", "Unit Price", "Total Price", "PDF pages"])
    for contractor in contractors:
        for item in contractor["items"]:
            source.append([
                contractor["name"], item["Item"], float(item["Qty"]), item["Unit"], item["Description"],
                float(item["Unit Price"]), float(item["Total Price"]), ", ".join(map(str, item["Pages"])),
            ])
            source.row_dimensions[source.max_row].height = 40
            source.cell(source.max_row, 6).number_format = '"$"#,##0.00'
            source.cell(source.max_row, 7).number_format = '"$"#,##0.00'
    # Reference style: contractor names above white headers, blue unit prices.
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = "E7"
    widths = [16, 14, 10, 54] + [22, 23] * len(contractors)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    edge = Side(style="thin", color="777777")
    grid = Border(left=edge, right=edge, top=edge, bottom=edge)
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A1"].font = Font(name="Arial", size=16, bold=True)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=sheet.max_column)
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=sheet.max_column)
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[4].height = 22
    sheet.row_dimensions[5].height = 44
    sheet.row_dimensions[6].height = 26
    for row in range(6, last_row + 1):
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, column)
            cell.border = grid
            if row == 6:
                cell.font = Font(name="Arial", size=11, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif column in (1, 2, 3):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif column >= 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
    for index, contractor in enumerate(contractors):
        unit_column, total_column = 5 + index * 2, 6 + index * 2
        sheet.merge_cells(start_row=5, start_column=unit_column, end_row=5, end_column=total_column)
        name_cell = sheet.cell(5, unit_column, contractor["name"])
        name_cell.font = Font(name="Arial", size=14, color="FF0000")
        name_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in range(first_row, last_row + 1):
            sheet.cell(row, unit_column).fill = PatternFill("solid", fgColor="DDF5F8")
        sheet.cell(base_row, unit_column).font = Font(name="Arial", size=11, bold=True)
        sheet.cell(base_row, unit_column).alignment = Alignment(horizontal="right", vertical="center")
        sheet.cell(base_row, total_column).font = Font(name="Arial", size=11, bold=True)
        sheet.cell(base_row, total_column).alignment = Alignment(horizontal="right", vertical="center")
        sheet.cell(base_row, total_column).border = Border(bottom=Side(style="double", color="555555"))
    for row in (base_row, calculated_row, difference_row, difference_row + 1):
        sheet.row_dimensions[row].height = 26
    sheet.row_dimensions[base_row].height = 30
    sheet.row_dimensions[difference_row + 1].height = 34
    # Keep validation visible below the base-bid row, outside the main grid.
    for row in (calculated_row, difference_row, difference_row + 1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    sheet.auto_filter.ref = f"A6:{get_column_letter(sheet.max_column)}{last_row}"
    style_sheet(source, [36, 16, 14, 10, 58, 20, 20, 14], 2)
    source.freeze_panes = "C3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:6"
    sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    workbook.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return str(path.resolve())


def save_master(summary: dict, contractors: list[dict], source_url: str, path: Path) -> str:
    from openpyxl import Workbook, load_workbook

    project = summary["project"]
    bid_date = None
    for pattern in ("%B %d, %Y", "%b %d, %Y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            bid_date = datetime.strptime(project["Bid Opening Date"], pattern)
            break
        except ValueError:
            continue
    if bid_date is None:
        raise ValueError(f"Unrecognized bid date: {project['Bid Opening Date']!r}.")
    headers = [
        "Airport", "CSJ", "Bid date", "Scope", "Contractor", "Item", "Description", "Qty", "Unit",
        "Unit price", "Total price", "Base bid", "PDF pages", "Source URL", "Status",
    ]
    existing = []
    if path.exists():
        previous = load_workbook(path, read_only=True, data_only=False)
        try:
            if previous.sheetnames != ["Master prices"]:
                raise ValueError("Master workbook has an unexpected layout; refusing to overwrite it.")
            rows = list(previous.active.values)
            if not rows or list(rows[0]) != headers:
                raise ValueError("Master workbook columns do not match; refusing to overwrite it.")
            for row in rows[1:]:
                # Replace this project's previous extraction; preserve other projects.
                if not (str(row[1]) == str(project["TxDOT CSJ#"]) and row[2] == bid_date):
                    existing.append(row)
        finally:
            previous.close()
    lowest = min(contractor["base_bid"] for contractor in contractors)
    winners = [contractor for contractor in contractors if contractor["base_bid"] == lowest]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master prices"
    sheet.append(headers)
    for row in existing:
        sheet.append(row)
    for contractor in winners:
        for item in contractor["items"]:
            sheet.append([
                project["Airport name"], project["TxDOT CSJ#"], bid_date, project["Project Scope"],
                contractor["name"], item["Item"], item["Description"], float(item["Qty"]), item["Unit"],
                float(item["Unit Price"]), float(item["Total Price"]), float(contractor["base_bid"]),
                ", ".join(map(str, item["Pages"])), source_url,
                "Apparent lowest (tie)" if len(winners) > 1 else "Apparent lowest",
            ])
    style_sheet(sheet, [34, 18, 16, 24, 38, 16, 58, 14, 10, 18, 20, 22, 14, 60, 25], 1)
    sheet.freeze_panes = "G2"
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 42
        sheet.cell(row, 3).number_format = "mm/dd/yyyy"
        sheet.cell(row, 8).number_format = "#,##0.###"
        for column in (10, 11, 12):
            sheet.cell(row, column).number_format = '"$"#,##0.00'
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return str(path.resolve())
